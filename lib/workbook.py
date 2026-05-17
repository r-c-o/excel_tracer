"""
lib/workbook.py
---------------
Thin wrapper around openpyxl that exposes:
  - Sheet names and dimensions
  - Raw formula strings (or values) for every cell
  - Named ranges
"""

from __future__ import annotations
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string


class Workbook:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        # data_only=False so we get formula strings, not cached values
        self._wb = openpyxl.load_workbook(self.path, data_only=False)
        self._wb_values = openpyxl.load_workbook(self.path, data_only=True)

    @property
    def sheet_names(self) -> list[str]:
        return self._wb.sheetnames

    def sheet(self, name: str):
        return self._wb[name]

    def sheet_values(self, name: str):
        return self._wb_values[name]

    def cell_formula(self, sheet_name: str, row: int, col: int) -> str | None:
        """Return formula string (with leading '=') or None if no formula."""
        ws = self._wb[sheet_name]
        cell = ws.cell(row=row, column=col)
        val = cell.value
        if isinstance(val, str) and val.startswith("="):
            return val
        return None

    def cell_value(self, sheet_name: str, row: int, col: int):
        """Return cached/computed value from the data_only workbook."""
        ws = self._wb_values[sheet_name]
        return ws.cell(row=row, column=col).value

    def iter_formula_cells(self, sheet_name: str):
        """Yield (row, col, formula_str) for every formula cell in a sheet."""
        ws = self._wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str) and val.startswith("="):
                    yield cell.row, cell.column, val

    def iter_all_cells(self, sheet_name: str):
        """Yield (row, col, value_or_formula) for every non-empty cell."""
        ws = self._wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    yield cell.row, cell.column, cell.value

    def dimensions(self, sheet_name: str) -> tuple[int, int, int, int]:
        """Return (min_row, min_col, max_row, max_col) for the used range."""
        ws = self._wb[sheet_name]
        return ws.min_row, ws.min_column, ws.max_row, ws.max_column

    def named_ranges(self) -> dict[str, str]:
        """Return {name: formula_string} for all workbook-level defined names."""
        result = {}
        for defn in self._wb.defined_names.definedName:
            result[defn.name] = defn.attr_text
        return result

    def col_letter(self, col_index: int) -> str:
        return get_column_letter(col_index)

    def col_index(self, letter: str) -> int:
        return column_index_from_string(letter.lstrip("$"))

    def header_row(self, sheet_name: str, header_row_num: int = 1) -> dict[str, int]:
        """Return {header_label: col_index} for a given header row."""
        ws = self._wb[sheet_name]
        headers = {}
        for cell in ws[header_row_num]:
            if cell.value is not None:
                headers[str(cell.value)] = cell.column
        return headers
